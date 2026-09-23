"""Bounded, deterministic adapters; source values never become executable code.

Rows in validation issues are one-based data rows (the header is not counted).
Mapping is explicit: ``amount`` means major currency units, never an inferred unit.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import re
import sys
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
import pyarrow.parquet as pq

from .models import (
    FORMATS, MAX_CELL_CHARS, MAX_COLUMNS, MAX_EDGES, MAX_FILE_BYTES,
    MAX_NODES, MAX_ROWS, Dataset, ImportFailure,
)

INT64_MAX = 2**63 - 1
MAX_AMOUNT_MINOR = INT64_MAX // 2  # Turnover counts each transfer at both ends.
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 1000
_OPTIONS = {"delimiter", "encoding", "sheet", "decimal", "date_format"}
_FIELDS = {
    "nodes": {"gid", "depth", "is_seed"},
    "edges": {"src", "dst", "amount", "n_tx", "depth", "currency"},
    "transactions": {"src", "dst", "amount", "date", "currency"},
}
_REQUIRED = {"nodes": {"gid"}, "edges": {"src", "dst", "amount"},
             "transactions": {"src", "dst", "amount"}}


def _fail(message: str, table: str = "file", row: int | None = None,
          field: str | None = None) -> None:
    raise ImportFailure(message, [{"table": table, "row": row, "field": field,
                                   "message": message}])


def _options(options: dict | None, extension: str) -> dict:
    if options is None:
        options = {}
    if not isinstance(options, dict) or set(options) - _OPTIONS:
        _fail("Неизвестные настройки формата.")
    result = {"encoding": "utf-8-sig", "delimiter": "\t" if extension == ".tsv" else ",",
              "sheet": None, "decimal": ".", "date_format": "iso", **options}
    if result["encoding"] not in ("utf-8", "utf-8-sig", "cp1251"):
        _fail("Кодировка должна быть utf-8, utf-8-sig или cp1251.")
    if result["delimiter"] not in (",", ";", "\t", "|"):
        _fail("Недопустимый разделитель; выберите запятую, точку с запятой, табуляцию или |.")
    if result["decimal"] not in (".", ","):
        _fail("Десятичный разделитель должен быть точкой или запятой.")
    if result["sheet"] is not None and (not isinstance(result["sheet"], str)
                                        or len(result["sheet"]) > MAX_CELL_CHARS):
        _fail("Имя листа должно быть строкой.")
    date_format = result["date_format"]
    if not isinstance(date_format, str) or not date_format or len(date_format) > 80:
        _fail("Укажите явный формат даты или iso.")
    if date_format != "iso":
        # No locale-dependent day/month names, timezone guessing or partial dates.
        directives = re.findall(r"%[A-Za-z%]", date_format)
        if (set(directives) - {"%Y", "%m", "%d", "%H", "%M", "%S", "%f", "%%"}
                or not {"%Y", "%m", "%d"}.issubset(directives)
                or re.search(r"%(?![YmdHMSf%])", date_format)):
            _fail("Формат даты должен явно включать %Y, %m, %d; часовой пояс не поддерживается.")
    return result


def normalize_config(config: dict) -> dict:
    defaults = {"profile": "generic", "currency": "KZT", "decimals": 0,
                "top_n": 30, "coverage": "unknown", "max_depth": None}
    if not isinstance(config, dict) or set(config) - set(defaults):
        _fail("Неизвестные параметры анализа.", "config")
    result = {**defaults, **config}
    if result["profile"] not in ("generic", "hackalem"):
        _fail("Неизвестный профиль анализа.", "config", field="profile")
    if result["profile"] == "hackalem":
        if "coverage" not in config:
            result["coverage"] = "outward"
        if "max_depth" not in config:
            result["max_depth"] = 4
    if not isinstance(result["currency"], str) or not re.fullmatch(r"[A-Z]{3}", result["currency"]):
        _fail("Валюта задаётся трёхбуквенным кодом в верхнем регистре.", "config", field="currency")
    for field, minimum, maximum in (("decimals", 0, 6), ("top_n", 1, 500)):
        value = result[field]
        if type(value) is not int or not minimum <= value <= maximum:
            _fail(f"Параметр {field} должен быть целым от {minimum} до {maximum}.", "config", field=field)
    if result["coverage"] not in ("unknown", "outward"):
        _fail("Неизвестные условия наблюдения.", "config", field="coverage")
    depth = result["max_depth"]
    if depth is not None and (type(depth) is not int or not 0 <= depth <= 1000):
        _fail("Предельная глубина должна быть целым от 0 до 1000.", "config", field="max_depth")
    if result["coverage"] == "outward" and depth is None:
        _fail("Для исходящего обхода укажите известную предельную глубину.", "config", field="max_depth")
    if result["coverage"] == "unknown" and depth is not None:
        _fail("Предельная глубина требует явно указанного исходящего обхода.", "config", field="max_depth")
    if result["profile"] == "hackalem":
        if (result["currency"] != "KZT" or result["decimals"] != 0 or result["top_n"] != 30
                or result["coverage"] != "outward" or result["max_depth"] != 4):
            _fail("Профиль HackAlem требует KZT, 0 десятичных знаков, top_n=30 и исходящий обход до глубины 4.", "config")
    return result


def _columns(columns: list[Any]) -> list[str]:
    if not columns or len(columns) > MAX_COLUMNS:
        _fail(f"Таблица должна содержать от 1 до {MAX_COLUMNS} колонок.")
    if any(not isinstance(name, str) or not name.strip() or len(name) > MAX_CELL_CHARS
           or any(ord(char) < 32 for char in name) for name in columns):
        _fail("Названия колонок должны быть непустыми строками без управляющих символов.")
    if len(set(columns)) != len(columns):
        _fail("Названия колонок должны быть уникальными.")
    return columns


def _cell(value: Any, row: int, column: str) -> Any:
    if isinstance(value, str):
        if len(value) > MAX_CELL_CHARS:
            _fail(f"Значение превышает {MAX_CELL_CHARS} символов.", row=row, field=column)
        return value
    if value is None or isinstance(value, (bool, int, float, Decimal, dt.date, dt.datetime)):
        # Missing Parquet values arrive as None; NaN/Infinity are not missing values.
        if isinstance(value, float) and not math.isfinite(value):
            _fail("NaN и бесконечные значения не поддерживаются.", row=row, field=column)
        if isinstance(value, Decimal) and not value.is_finite():
            _fail("NaN и бесконечные значения не поддерживаются.", row=row, field=column)
        return value
    _fail("Поддерживаются только плоские скалярные значения; вложенные объекты запрещены.", row=row, field=column)


def _append(rows: list[dict], columns: list[str], values: list[Any], decoded_bytes: int) -> int:
    number = len(rows) + 1
    if number > MAX_ROWS:
        _fail(f"Превышен предел {MAX_ROWS} строк.", row=number)
    if len(values) != len(columns):
        _fail("Количество значений не совпадает с заголовком таблицы.", row=number)
    record = {name: _cell(value, number, name) for name, value in zip(columns, values)}
    # Parquet metadata measures encoded pages: dictionary values can expand
    # many times when materialized. Bound accumulated Python values as well.
    # Conservatively count repeated scalars even if a reader shares them.
    decoded_bytes += sys.getsizeof(record) + sum(sys.getsizeof(value) for value in record.values())
    if decoded_bytes > MAX_UNPACKED_BYTES:
        _fail("Превышен лимит объёма декодированной таблицы (100 MiB).", row=number)
    rows.append(record)
    return decoded_bytes


def _unique_object(pairs: list[tuple[str, Any]]) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            _fail("JSON содержит повторяющийся ключ объекта.")
        result[key] = value
    return result


def _bad_constant(value: str) -> None:
    _fail("JSON не должен содержать NaN или Infinity.")


def _read_json(text: str) -> Any:
    return json.loads(text, parse_float=Decimal, parse_constant=_bad_constant,
                      object_pairs_hook=_unique_object)


def _xlsx_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_MEMBERS or sum(item.file_size for item in members) > MAX_UNPACKED_BYTES:
            _fail("XLSX превышает предел распаковки 100 MiB или 1000 элементов.")
        if any(item.flag_bits & 1 for item in members):
            _fail("Зашифрованные XLSX не поддерживаются.")
        if any("vbaproject" in item.filename.lower() for item in members):
            _fail("Книги с макросами не поддерживаются.")
        # Explicitly reject XML entities/DTDs even if optional XML parsers change.
        for member in members:
            if member.filename.lower().endswith((".xml", ".rels")):
                with archive.open(member) as xml:
                    carry = b""
                    while chunk := xml.read(65536):
                        value = (carry + chunk).upper()
                        if b"<!DOCTYPE" in value or b"<!ENTITY" in value:
                            _fail("XML-объявления DTD и сущностей в XLSX запрещены.")
                        carry = value[-16:]


def _read(path: Path, filename: str, options: dict | None = None) -> tuple[list[str], list[dict], list[str], dict]:
    if not isinstance(filename, str) or not filename or len(filename) > MAX_CELL_CHARS:
        _fail("Некорректное имя файла.")
    extension = Path(filename).suffix.lower()
    if extension not in FORMATS:
        _fail("Поддерживаются CSV, TSV, XLSX, JSON, JSONL и Parquet.")
    settings = _options(options, extension)
    rows: list[dict] = []
    decoded_bytes = 0
    sheets: list[str] = []
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            _fail("Размер файла превышает 20 MiB.")
        if extension in (".csv", ".tsv"):
            with path.open("r", encoding=settings["encoding"], newline="") as stream:
                reader = csv.reader(stream, delimiter=settings["delimiter"], strict=True)
                columns = _columns(next(reader, []))
                for values in reader:
                    decoded_bytes = _append(rows, columns, values, decoded_bytes)
        elif extension in (".json", ".jsonl"):
            with path.open("r", encoding=settings["encoding"]) as stream:
                if extension == ".json":
                    documents = _read_json(stream.read())
                    if not isinstance(documents, list):
                        _fail("JSON должен быть массивом плоских объектов.")
                else:
                    documents = []
                    for line in stream:
                        if line.strip():
                            documents.append(_read_json(line))
                            if len(documents) > MAX_ROWS:
                                _fail(f"Превышен предел {MAX_ROWS} строк.")
            if not documents or len(documents) > MAX_ROWS:
                _fail(f"Таблица должна содержать от 1 до {MAX_ROWS} строк.")
            if any(not isinstance(row, dict) for row in documents):
                _fail("Каждая строка JSON должна быть плоским объектом.")
            columns = _columns(list(dict.fromkeys(key for row in documents for key in row)))
            for document in documents:
                decoded_bytes = _append(rows, columns, [document.get(name) for name in columns], decoded_bytes)
        elif extension == ".parquet":
            # Own the descriptor even when the ParquetFile constructor fails.
            with path.open("rb") as parquet_stream:
                parquet = pq.ParquetFile(parquet_stream)
                try:
                    if parquet.metadata.num_rows > MAX_ROWS:
                        _fail(f"Превышен предел {MAX_ROWS} строк.")
                    unpacked = sum(parquet.metadata.row_group(i).total_byte_size
                                   for i in range(parquet.metadata.num_row_groups))
                    if unpacked > MAX_UNPACKED_BYTES:
                        _fail("Распакованный Parquet превышает 100 MiB.")
                    columns = _columns(parquet.schema_arrow.names)
                    for batch in parquet.iter_batches(batch_size=64):
                        for document in batch.to_pylist():
                            decoded_bytes = _append(rows, columns, [document[name] for name in columns], decoded_bytes)
                finally:
                    parquet.close()
        else:
            from openpyxl import load_workbook
            _xlsx_archive(path)
            with path.open("rb") as workbook_stream:
                workbook = load_workbook(workbook_stream, read_only=True, data_only=False, keep_links=False)
                try:
                    sheets = workbook.sheetnames
                    selected = settings["sheet"] if settings["sheet"] is not None else sheets[0]
                    if selected not in sheets:
                        _fail("Выбранный лист отсутствует в книге.", field="sheet")
                    sheet = workbook[selected]
                    if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                        _fail("Размер листа превышает предел строк или колонок.")
                    iterator = sheet.iter_rows()
                    header = next(iterator, ())
                    if any(cell.data_type == "f" for cell in header):
                        _fail("Формулы в заголовке XLSX запрещены.")
                    columns = _columns([cell.value for cell in header])
                    for cells in iterator:
                        values = []
                        for cell in cells:
                            if cell.data_type == "f":
                                _fail("Формулы в XLSX запрещены; сохраните их результаты как значения.",
                                      row=len(rows) + 1, field=columns[len(values)])
                            value = cell.value
                            unsafe_integer = type(value) is int and abs(value) > 999999999999999
                            unsafe_float = (type(value) is float and math.isfinite(value)
                                            and len(Decimal(str(value)).normalize().as_tuple().digits) > 15)
                            if unsafe_integer or unsafe_float:
                                _fail("Число XLSX длиннее 15 значащих цифр неоднозначно; храните идентификаторы и точные суммы как текст.",
                                      row=len(rows) + 1, field=columns[len(values)])
                            values.append(value)
                        decoded_bytes = _append(rows, columns, values, decoded_bytes)
                finally:
                    workbook.close()
    except ImportFailure:
        raise
    except Exception as error:
        # Do not expose library exception messages, which can contain raw values/paths.
        raise ImportFailure("Файл повреждён или не соответствует выбранному формату и настройкам.",
                            [{"table": "file", "row": None, "field": None,
                              "message": "Не удалось прочитать таблицу; проверьте формат, кодировку и разделитель."}]) from error
    if not rows:
        _fail("Таблица не содержит строк данных.")
    return columns, rows, sheets, settings


def _preview(value: Any) -> Any:
    # Integers are strings in preview, too: JS must not round 64-bit identifiers.
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, (Decimal, int)) and not isinstance(value, bool):
        return str(value)
    return value


def inspect_file(path: Path, filename: str, options: dict | None = None) -> dict:
    columns, rows, sheets, _ = _read(Path(path), filename, options)
    return {"columns": columns, "preview": [{key: _preview(value) for key, value in row.items()}
                                            for row in rows[:5]],
            "rows": len(rows), "sheets": sheets, "warnings": []}


def _missing(value: Any) -> bool:
    return value is None or value == ""


def _identifier(value: Any) -> str:
    if isinstance(value, bool) or value is None:
        raise ValueError("ID должен быть непустой строкой или целым числом.")
    if isinstance(value, int):
        return str(value)
    if not isinstance(value, str) or not value or value != value.strip():
        raise ValueError("ID должен быть непустой строкой без внешних пробелов или целым числом; дробные ID запрещены.")
    if any(ord(char) < 32 or ord(char) == 127 for char in value):
        raise ValueError("ID не должен содержать управляющие символы.")
    if re.fullmatch(r"[+-]?(?:\d+\.\d*|\.\d+)(?:[eE][+-]?\d+)?", value):
        raise ValueError("Дробная запись ID не поддерживается; используйте исходный текстовый идентификатор.")
    return value


def _integer(value: Any, minimum: int = 0, maximum: int = INT64_MAX) -> int | None:
    if _missing(value):
        return None
    if isinstance(value, bool) or (not isinstance(value, (int, str))):
        raise ValueError("Требуется целое число.")
    if isinstance(value, str) and not re.fullmatch(r"[+-]?\d+", value):
        raise ValueError("Требуется целое число без округления.")
    number = int(value)
    if not minimum <= number <= maximum:
        raise ValueError(f"Целое число должно находиться в диапазоне {minimum}–{maximum}.")
    return number


def _boolean(value: Any) -> bool | None:
    if _missing(value):
        return None
    if type(value) is bool:
        return value
    if type(value) is int and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.lower() in ("true", "false", "0", "1"):
        return value.lower() in ("true", "1")
    raise ValueError("Признак seed должен быть true/false, 1/0 или пустым.")


def _money(value: Any, decimals: int, separator: str) -> int:
    if isinstance(value, bool) or value is None or isinstance(value, (dt.date, dt.datetime)):
        raise ValueError("Сумма должна быть конечным положительным числом.")
    if isinstance(value, str):
        if separator == ",":
            if "." in value:
                raise ValueError("Сумма не соответствует выбранному десятичному разделителю.")
            value = value.replace(",", ".")
        if not re.fullmatch(r"[+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value):
            raise ValueError("Сумма должна быть положительным числом без разделителей тысяч.")
    try:
        amount = Decimal(str(value))
        if not amount.is_finite() or amount <= 0 or amount.adjusted() > 18:
            raise ValueError("Сумма должна быть конечным положительным числом в допустимом диапазоне.")
        # Decimal context must not round a long coefficient during multiplication.
        sign, digits, exponent = amount.as_tuple()
        scaled_exponent = exponent + decimals
        if scaled_exponent < 0:
            if -scaled_exponent > len(digits) or any(digits[scaled_exponent:]):
                raise ValueError("Точность суммы превышает выбранное число десятичных знаков; округление не выполняется.")
            digits = digits[:scaled_exponent]
            scaled_exponent = 0
        if len(digits) + scaled_exponent > 19:
            raise ValueError("Сумма превышает безопасный предел int64 для оборотов первой версии.")
        scaled = int("".join(str(digit) for digit in digits) or "0") * 10**scaled_exponent
        if not 0 < scaled <= MAX_AMOUNT_MINOR:
            raise ValueError("Сумма превышает безопасный предел int64 для оборотов первой версии.")
        return scaled
    except (InvalidOperation, OverflowError) as error:
        raise ValueError("Сумма имеет недопустимый числовой формат.") from error


def _date(value: Any, date_format: str) -> dt.datetime | None:
    if _missing(value):
        return None
    if isinstance(value, dt.datetime):
        result = value
    elif isinstance(value, dt.date):
        result = dt.datetime.combine(value, dt.time())
    elif isinstance(value, str):
        try:
            if date_format == "iso":
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?", value):
                    raise ValueError()
                result = dt.datetime.fromisoformat(value)
            else:
                result = dt.datetime.strptime(value, date_format)
        except ValueError as error:
            raise ValueError("Дата не соответствует выбранному формату или содержит часовой пояс.") from error
    else:
        raise ValueError("Дата должна быть текстовой датой или датой XLSX/Parquet, не числом.")
    if result.tzinfo is not None:
        raise ValueError("Даты с часовым поясом требуют предварительного явного приведения; зона не угадывается.")
    return result


def _normalize(kind: str, rows: list[dict], columns: list[str], mapping: dict,
               settings: dict, config: dict) -> list[dict]:
    if not isinstance(mapping, dict) or set(mapping) - _FIELDS[kind]:
        _fail("Неизвестные поля сопоставления.", kind)
    if not _REQUIRED[kind].issubset(mapping):
        _fail("Сопоставьте все обязательные поля таблицы.", kind)
    if any(not isinstance(source, str) or source not in columns for source in mapping.values()):
        _fail("Сопоставление ссылается на отсутствующую колонку.", kind)
    if len(set(mapping.values())) != len(mapping):
        _fail("Одна исходная колонка не может иметь несколько назначений.", kind)
    result = []
    issues = []
    for number, source in enumerate(rows, start=1):
        record = {}
        for field in sorted(_FIELDS[kind]):
            value = source[mapping[field]] if field in mapping else None
            try:
                if field in ("gid", "src", "dst"):
                    record[field] = _identifier(value)
                elif field == "amount":
                    record["amount_minor"] = _money(value, config["decimals"], settings["decimal"])
                elif field in ("depth", "n_tx"):
                    record[field] = _integer(value, minimum=1 if field == "n_tx" else 0,
                                             maximum=INT64_MAX if field == "n_tx" else 1000)
                elif field == "is_seed":
                    record[field] = _boolean(value)
                elif field == "date":
                    record[field] = _date(value, settings["date_format"])
                elif field == "currency" and field in mapping:
                    if value != config["currency"]:
                        raise ValueError("Валюта строки отсутствует или отличается от валюты анализа; конвертации нет.")
            except (ValueError, TypeError) as error:
                issues.append({"table": kind, "row": number, "field": field, "message": str(error)})
                if len(issues) >= 100:
                    raise ImportFailure("Исправьте ошибки входных данных (показаны первые 100).", issues)
        result.append(record)
    if issues:
        raise ImportFailure("Исправьте ошибки входных данных.", issues)
    return result


def prepare_dataset(sources: list[dict], config: dict) -> Dataset:
    """Normalize a transaction file or coherent tables without invented metadata."""
    config = normalize_config(config)
    if not isinstance(sources, list) or not 1 <= len(sources) <= 3:
        _fail("Нужны от одной до трёх входных таблиц.")
    tables = {}
    warnings = []
    for source in sources:
        if not isinstance(source, dict) or set(source) - {"kind", "path", "filename", "mapping", "options"}:
            _fail("Некорректное описание источника.")
        kind = source.get("kind")
        if not isinstance(kind, str) or kind not in _FIELDS or kind in tables:
            _fail("Назначение таблицы должно быть nodes, edges или transactions, без повторов.")
        if not {"path", "filename", "mapping"}.issubset(source):
            _fail("У источника отсутствует путь, имя или сопоставление полей.", kind)
        try:
            if not isinstance(source["path"], (str, Path)):
                _fail("Путь источника должен быть назначен сервером.", kind)
            columns, rows, _, settings = _read(Path(source["path"]), source["filename"], source.get("options"))
        except ImportFailure as error:
            for issue in error.issues:
                issue["table"] = kind
            raise
        tables[kind] = _normalize(kind, rows, columns, source["mapping"], settings, config)
    if "transactions" not in tables and "edges" not in tables:
        _fail("Нужна таблица транзакций или направленных рёбер.")
    if config["profile"] == "hackalem" and set(tables) != set(_FIELDS):
        _fail("Для HackAlem нужны все три таблицы: nodes, edges и transactions.")
    transactions = tables.get("transactions")
    if transactions is not None:
        aggregate = {}
        seen_transactions = Counter()
        for number, row in enumerate(transactions, start=1):
            if row["src"] == row["dst"]:
                _fail("Петли (перевод узла самому себе) не поддерживаются.", "transactions", number, "dst")
            pair = (row["src"], row["dst"])
            item = aggregate.setdefault(pair, {"src": pair[0], "dst": pair[1], "amount_minor": 0,
                                               "n_tx": 0, "depth": None})
            item["amount_minor"] += row["amount_minor"]
            item["n_tx"] += 1
            seen_transactions[(row["src"], row["dst"], row["amount_minor"], row["date"])] += 1
        duplicate_count = sum(count - 1 for count in seen_transactions.values())
        if duplicate_count:
            warnings.append({"table": "transactions", "row": None, "field": None,
                             "message": f"Совпадающих по импортированным полям строк: {duplicate_count}. Они сохранены; совпадение не доказывает дубль."})
    else:
        aggregate = None
    edges = tables.get("edges")
    if edges is None:
        edges = list(aggregate.values())
    edge_pairs = set()
    for number, edge in enumerate(edges, start=1):
        pair = (edge["src"], edge["dst"])
        if pair[0] == pair[1]:
            _fail("Петли (перевод узла самому себе) не поддерживаются.", "edges", number, "dst")
        if pair in edge_pairs:
            _fail("Повторная направленная пара рёбер; предварительно согласуйте агрегацию.", "edges", number, "src")
        edge_pairs.add(pair)
        if aggregate is not None:
            expected = aggregate.get(pair)
            if expected is None or expected["amount_minor"] != edge["amount_minor"]:
                _fail("Сумма ребра не совпадает с суммой транзакций этой пары.", "edges", number, "amount")
            if edge["n_tx"] is not None and edge["n_tx"] != expected["n_tx"]:
                _fail("Число операций ребра не совпадает с числом транзакций.", "edges", number, "n_tx")
            if edge["n_tx"] is None:
                edge["n_tx"] = expected["n_tx"]
    if aggregate is not None and edge_pairs != set(aggregate):
        _fail("Набор пар в рёбрах не совпадает с транзакциями.", "edges")
    if len(edges) > MAX_EDGES:
        _fail(f"Превышен предел {MAX_EDGES} направленных рёбер.", "edges")
    total = sum(row["amount_minor"] for row in edges)
    if total > MAX_AMOUNT_MINOR:
        _fail("Общая сумма превышает (2^63−1)/2 минимальных единиц — безопасный предел int64 для оборотов первой версии.", "edges", field="amount")
    if sum(row["n_tx"] or 0 for row in edges) > INT64_MAX:
        _fail("Общее число операций превышает int64.", "edges", field="n_tx")
    nodes = tables.get("nodes")
    endpoints = {row[field] for row in edges for field in ("src", "dst")}
    if nodes is None:
        nodes = [{"gid": gid, "depth": None, "is_seed": None} for gid in sorted(endpoints)]
    node_ids = set()
    for number, node in enumerate(nodes, start=1):
        if node["gid"] in node_ids:
            _fail("Идентификатор узла повторяется.", "nodes", number, "gid")
        node_ids.add(node["gid"])
        if config["coverage"] == "outward" and node["depth"] is not None and node["depth"] > config["max_depth"]:
            _fail("Глубина узла превышает явно заявленную границу наблюдения.", "nodes", number, "depth")
        if (config["coverage"] == "outward" and node["depth"] is not None and node["is_seed"] is not None
                and node["is_seed"] != (node["depth"] == 0)):
            _fail("Признак seed противоречит глубине исходящего обхода (seed соответствует глубине 0).", "nodes", number, "is_seed")
    if not endpoints.issubset(node_ids):
        for number, edge in enumerate(edges, start=1):
            for field in ("src", "dst"):
                if edge[field] not in node_ids:
                    _fail("Узел ребра отсутствует в явной таблице nodes.", "edges", number, field)
    if len(nodes) > MAX_NODES:
        _fail(f"Превышен предел {MAX_NODES} узлов.", "nodes")
    depths = {node["gid"]: node["depth"] for node in nodes}
    for number, edge in enumerate(edges, start=1):
        if (edge["depth"] is not None and depths[edge["dst"]] is not None
                and edge["depth"] != depths[edge["dst"]]):
            _fail("Глубина ребра не совпадает с глубиной его получателя.", "edges", number, "depth")
        if config["coverage"] == "outward" and depths[edge["src"]] == config["max_depth"]:
            _fail("У узла на заявленной границе есть исходящее ребро: условия обрыва обхода противоречат данным.", "edges", number, "src")
    if any(row["is_seed"] is None for row in nodes):
        warnings.append({"table": "nodes", "row": None, "field": "is_seed",
                         "message": "Признак seed известен не для всех узлов; зависимые выводы ограничены."})
    if any(row["depth"] is None for row in nodes):
        warnings.append({"table": "nodes", "row": None, "field": "depth",
                         "message": "Глубина обхода известна не для всех узлов; граница не выводится из максимума данных."})
    if transactions is None:
        warnings.append({"table": "transactions", "row": None, "field": None,
                         "message": "Отдельные транзакции не переданы; даты и временные паттерны недоступны."})
    elif any(row["date"] is None for row in transactions):
        warnings.append({"table": "transactions", "row": None, "field": "date",
                         "message": "Даты известны не для всех транзакций; временной анализ ограничен."})
    nodes_frame = pd.DataFrame(nodes, columns=["gid", "depth", "is_seed"])
    nodes_frame["depth"] = pd.array(nodes_frame["depth"], dtype="Int64")
    nodes_frame["is_seed"] = pd.array(nodes_frame["is_seed"], dtype="boolean")
    edges_frame = pd.DataFrame(edges, columns=["src", "dst", "amount_minor", "n_tx", "depth"])
    edges_frame["amount_minor"] = edges_frame["amount_minor"].astype("int64")
    for column in ("n_tx", "depth"):
        edges_frame[column] = pd.array(edges_frame[column], dtype="Int64")
    tx_frame = None
    if transactions is not None:
        tx_frame = pd.DataFrame(transactions, columns=["src", "dst", "amount_minor", "date"])
        tx_frame["amount_minor"] = tx_frame["amount_minor"].astype("int64")
        tx_frame["date"] = pd.Series([row["date"] for row in transactions], dtype=object)
    return Dataset(nodes_frame, edges_frame, tx_frame, dict(config), warnings)
